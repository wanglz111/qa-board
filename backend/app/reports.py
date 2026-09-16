from __future__ import annotations

import csv
import io
import re
from datetime import datetime
from typing import Any
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Response
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.auth import require_admin
from app.db import get_db
from app.models import Attempt, CaseReferenceLink, Group, GroupCase

router = APIRouter(prefix="/api", dependencies=[Depends(require_admin)])

# Spreadsheet consumers execute any cell that starts with these characters, so
# exported text is defused with a leading apostrophe instead of being trusted.
FORMULA_PREFIXES = ("=", "+", "-", "@")
# A formula can hide behind leading whitespace or a control character, and some
# consumers strip those before evaluating the cell.
IGNORED_LEADING_CHARACTERS = " \t\r\n"
CSV_MEDIA_TYPE = "text/csv; charset=utf-8"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADERS = (
    "source_file",
    "source_version",
    "code",
    "position",
    "title",
    "module",
    "layer",
    "priority",
    "result",
    "note",
    "executed_at",
    "attempt_label",
    "history_count",
    "screenshot_count",
    "reference_image_count",
    "source",
)


def _safe(value: Any) -> Any:
    if isinstance(value, str):
        candidate = value.lstrip(IGNORED_LEADING_CHARACTERS)
        if candidate.startswith(FORMULA_PREFIXES):
            return f"'{value}"
    return value


def _timestamp(value: datetime | None) -> str | None:
    return value.isoformat() if value is not None else None


def _group_or_404(db: Session, group_id: UUID) -> Group:
    group = db.get(Group, group_id)
    if group is None:
        raise HTTPException(status_code=404, detail="Group not found")
    return group


def _report_rows(db: Session, group: Group) -> list[dict[str, Any]]:
    cases = db.scalars(
        select(GroupCase).where(GroupCase.group_id == group.id).order_by(GroupCase.position)
    ).all()
    reference_counts = _reference_counts(db, group)
    attempts = db.scalars(
        select(Attempt)
        .join(GroupCase, Attempt.group_case_id == GroupCase.id)
        .where(GroupCase.group_id == group.id, Attempt.state == "committed")
        .order_by(Attempt.group_case_id, Attempt.sequence)
    ).all()
    history: dict[UUID, list[Attempt]] = {}
    for attempt in attempts:
        history.setdefault(attempt.group_case_id, []).append(attempt)

    rows: list[dict[str, Any]] = []
    for group_case in cases:
        case_history = history.get(group_case.id, [])
        latest = case_history[-1] if case_history else None
        rows.append(
            {
                "source_file": group.source_name,
                "source_version": group.source_version,
                "code": group_case.code,
                "position": group_case.position,
                "title": group_case.title,
                "module": group_case.module,
                "layer": group_case.layer,
                "priority": group_case.priority,
                "result": latest.result if latest else None,
                "note": latest.note if latest else None,
                "executed_at": _timestamp(latest.created_at) if latest else None,
                "attempt_label": latest.label if latest else None,
                "history_count": len(case_history),
                "screenshot_count": sum(len(attempt.screenshots) for attempt in case_history),
                "reference_image_count": reference_counts.get(group_case.id, 0),
                "source": latest.source if latest else None,
            }
        )
    return rows


def _reference_counts(db: Session, group: Group) -> dict[UUID, int]:
    rows = db.execute(
        select(CaseReferenceLink.group_case_id, func.count(CaseReferenceLink.id))
        .join(GroupCase, CaseReferenceLink.group_case_id == GroupCase.id)
        .where(GroupCase.group_id == group.id)
        .group_by(CaseReferenceLink.group_case_id)
    ).all()
    return dict(rows)


def _filename(group: Group, suffix: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9._-]+", "-", group.short_code).strip("-") or "group"
    return f"testdeck-{stem}.{suffix}"


def _attachment_headers(group: Group, suffix: str) -> dict[str, str]:
    return {
        "Content-Disposition": f'attachment; filename="{_filename(group, suffix)}"',
        "Cache-Control": "private, no-store",
    }


@router.get("/groups/{group_id}/reports.csv")
def group_report_csv(group_id: UUID, db: Session = Depends(get_db)) -> Response:
    group = _group_or_404(db, group_id)
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(HEADERS))
    writer.writeheader()
    for row in _report_rows(db, group):
        writer.writerow({key: _safe(value) for key, value in row.items()})
    return Response(
        content=buffer.getvalue(),
        media_type=CSV_MEDIA_TYPE,
        headers=_attachment_headers(group, "csv"),
    )


@router.get("/groups/{group_id}/reports.xlsx")
def group_report_xlsx(group_id: UUID, db: Session = Depends(get_db)) -> Response:
    group = _group_or_404(db, group_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "group-report"
    sheet.append([header.replace("_", " ") for header in HEADERS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in _report_rows(db, group):
        sheet.append([_safe(row[header]) for header in HEADERS])
    for index, header in enumerate(HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(
            12, min(40, len(header) + 6)
        )
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type=XLSX_MEDIA_TYPE,
        headers=_attachment_headers(group, "xlsx"),
    )
