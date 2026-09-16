import csv
import io
from uuid import uuid4

from openpyxl import load_workbook

from tests.casebook_fixture import casebook_zip


def _rows(response) -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(response.text)))


def test_export_is_group_scoped_and_blocks_spreadsheet_formula(
    authenticated_client, imported_group, add_case
):
    add_case(imported_group.id, code="X-001", title='=HYPERLINK("https://unsafe.test")')

    output = authenticated_client.get(f"/api/groups/{imported_group.id}/reports.csv")

    assert output.status_code == 200
    assert output.headers["content-disposition"].startswith("attachment;")
    assert output.headers["cache-control"] == "private, no-store"
    titles = [row["title"] for row in _rows(output)]
    assert titles == ["管理员登录", '\'=HYPERLINK("https://unsafe.test")']
    assert "'=HYPERLINK" in titles[-1]
    assert output.headers["content-disposition"].endswith('.csv"')


def test_csv_contains_source_metadata_latest_attempt_and_counts(
    authenticated_client, imported_group, add_case, valid_png, upload_dir
):
    add_case(imported_group.id, code="B-002", title="+SUM(1,1)")
    attempts_url = f"/api/groups/{imported_group.id}/cases/B-001/attempts"
    first = authenticated_client.post(
        attempts_url,
        json={
            "result": "不通过",
            "note": "@cmd failing note",
            "idempotency_key": "report-1",
        },
    )
    assert first.status_code == 201
    authenticated_client.post(
        f"/api/attempts/{first.json()['id']}/screenshots",
        files={"image": ("shot.png", valid_png, "image/png")},
    )
    second = authenticated_client.post(
        attempts_url,
        json={"result": "通过", "idempotency_key": "report-2"},
    )
    assert second.status_code == 201

    rows = _rows(authenticated_client.get(f"/api/groups/{imported_group.id}/reports.csv"))

    assert len(rows) == 2
    first_row = rows[0]
    assert first_row["source_file"] == imported_group.source_name
    assert first_row["source_version"] == imported_group.source_version
    assert first_row["code"] == "B-001"
    assert first_row["result"] == "通过"
    assert first_row["attempt_label"] == second.json()["label"]
    assert first_row["history_count"] == "2"
    assert first_row["screenshot_count"] == "1"
    assert first_row["executed_at"]
    # Provenance closes the row, and a locally executed attempt says so.
    assert list(first_row)[-1] == "source"
    assert first_row["source"] == "execution"
    assert rows[1]["title"] == "'+SUM(1,1)"
    assert rows[1]["result"] == ""
    assert rows[1]["history_count"] == "0"
    # Private screenshots stay behind the authenticated route: no URL is exported.
    assert "/api/screenshots" not in authenticated_client.get(
        f"/api/groups/{imported_group.id}/reports.csv"
    ).text


def test_xlsx_keeps_malicious_text_literal(
    authenticated_client, imported_group, add_case
):
    add_case(imported_group.id, code="X-001", title='=HYPERLINK("https://unsafe.test")')
    add_case(imported_group.id, code="X-002", title="@SUM(A1)")

    output = authenticated_client.get(f"/api/groups/{imported_group.id}/reports.xlsx")

    assert output.status_code == 200
    workbook = load_workbook(io.BytesIO(output.content))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    assert headers[headers.index("title")] == "title"
    assert headers[-1] == "source"
    titles = [sheet.cell(row=index, column=headers.index("title") + 1).value for index in (3, 4)]
    assert titles == ["'=HYPERLINK(\"https://unsafe.test\")", "'@SUM(A1)"]
    assert sheet.cell(row=2, column=headers.index("source file") + 1).value == "0918.csv"


def test_export_requires_a_session_and_rejects_unknown_group(
    authenticated_client, anonymous_client, imported_group
):
    assert (
        anonymous_client.get(f"/api/groups/{imported_group.id}/reports.csv").status_code
        == 401
    )
    missing = authenticated_client.get(f"/api/groups/{uuid4()}/reports.csv")
    assert missing.status_code == 404


def test_export_defuses_formulas_hidden_behind_leading_whitespace(
    authenticated_client, imported_group, add_case
):
    add_case(imported_group.id, code="X-001", title="\t=HYPERLINK(\"https://unsafe.test\")")
    add_case(imported_group.id, code="X-002", title=" =1+1")
    add_case(imported_group.id, code="X-003", title="\r\n-2+3")

    csv_rows = _rows(authenticated_client.get(f"/api/groups/{imported_group.id}/reports.csv"))
    assert [row["title"] for row in csv_rows] == [
        "管理员登录",
        "'\t=HYPERLINK(\"https://unsafe.test\")",
        "' =1+1",
        "'\r\n-2+3",
    ]

    workbook = load_workbook(
        io.BytesIO(authenticated_client.get(f"/api/groups/{imported_group.id}/reports.xlsx").content)
    )
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    titles = [
        sheet.cell(row=index, column=headers.index("title") + 1).value for index in (3, 4, 5)
    ]
    # openpyxl normalises the CRLF pair to a single newline on write.
    assert titles == ["'\t=HYPERLINK(\"https://unsafe.test\")", "' =1+1", "'\n-2+3"]


def test_report_counts_reference_images_separately(authenticated_client, upload_dir):
    preview = authenticated_client.post(
        "/api/import/preview",
        files={"file": ("casebook.zip", casebook_zip(), "application/zip")},
    )
    group_id = authenticated_client.post(
        "/api/import/confirm",
        json={"ticket_id": preview.json()["ticket_id"], "name": "Odyssey"},
    ).json()["id"]

    rows = _rows(authenticated_client.get(f"/api/groups/{group_id}/reports.csv"))

    # C-05 references two images; C-11 references one locator.
    assert rows[0]["reference_image_count"] == "2"
    assert rows[1]["reference_image_count"] == "1"
    assert rows[0]["screenshot_count"] == "0"
    assert rows[0]["source"] == ""
    assert list(rows[0])[-1] == "source"
